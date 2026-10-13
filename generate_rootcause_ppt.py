"""
生成 Agent 运行根因定位方案 PPT (一页)
布局: 标题 + 三栏 (因果图层级 | 分析流程 | 推理方法) + 底部核心思路
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ===== Color palette =====
NAVY       = RGBColor(0x1F, 0x3A, 0x68)
BLUE       = RGBColor(0x3B, 0x82, 0xF6)
LIGHT_BLUE = RGBColor(0xDB, 0xEA, 0xFE)
GREEN      = RGBColor(0x10, 0xB9, 0x81)
RED        = RGBColor(0xEF, 0x44, 0x44)
YELLOW     = RGBColor(0xF5, 0x9E, 0x0B)
PURPLE     = RGBColor(0x8B, 0x5C, 0xF6)
GRAY       = RGBColor(0x6B, 0x72, 0x80)
LIGHT_GRAY = RGBColor(0xF3, 0xF4, 0xF6)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK       = RGBColor(0x1F, 0x29, 0x37)

# ===== Init presentation (16:9) =====
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

def box(left, top, w, h, text, fill, fc=DARK, fs=11, bold=False, line_color=None):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(left), Inches(top), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = fill
    if line_color:
        b.line.color.rgb = line_color; b.line.width = Pt(1.5)
    else:
        b.line.fill.background()
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    tf.margin_top  = Inches(0.04); tf.margin_bottom = Inches(0.04)
    tf.text = text
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(fs); p.font.bold = bold; p.font.color.rgb = fc
    return b

def arrow(left, top, w, h, color=BLUE):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,
                               Inches(left), Inches(top), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background()
    return a

def txt(left, top, w, h, text, fc=DARK, fs=14, bold=False, italic=False, align=PP_ALIGN.LEFT):
    t = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = t.text_frame; tf.word_wrap = True
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = align; p.font.size = Pt(fs); p.font.bold = bold
    p.font.italic = italic; p.font.color.rgb = fc
    return t

# ===== Title =====
txt(0.5, 0.20, 12.33, 0.55, "Agent 运行根因定位方案", NAVY, 28, bold=True)
txt(0.5, 0.78, 12.33, 0.3,  "基于因果图的回溯分析  ·  5Why + CAE 推理方法", GRAY, 12, italic=True)
line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.12), Inches(12.33), Inches(0.02))
line.fill.solid(); line.fill.fore_color.rgb = NAVY; line.line.fill.background()

# ===== Column 1: 因果图节点层级 =====
txt(0.5, 1.28, 6.0, 0.4, "① 因果图节点层级  (Causal Graph)", NAVY, 15, bold=True)

box(1.7, 1.82, 3.0, 0.45, "🎯  用户需求 / 目标",      GREEN,  WHITE, 12, True)
arrow(2.95, 2.32, 0.5, 0.30, GREEN)
box(1.7, 2.67, 3.0, 0.45, "⚙️  执行 Skill  (L1)",     NAVY,   WHITE, 12, True)
arrow(2.95, 3.17, 0.5, 0.30, NAVY)
box(1.7, 3.52, 3.0, 0.45, "📋  Task + 产出 .md  (L2)", BLUE,   WHITE, 12, True)
arrow(2.95, 4.02, 0.5, 0.30, BLUE)

# L3: 4 parallel nodes
box(0.30, 4.37, 1.50, 0.60, "🔧 tool_call\nbash / read / write / edit", LIGHT_BLUE, NAVY,   9,  False)
box(1.90, 4.37, 1.25, 0.60, "💭 thought",  LIGHT_GRAY, NAVY,   10, True)
box(3.25, 4.37, 1.25, 0.60, "✅ decision", PURPLE,     WHITE,  10, True)
box(4.60, 4.37, 1.25, 0.60, "❌ error",    RED,        WHITE,  10, True)

arrow(2.95, 5.02, 0.5, 0.30, GRAY)
box(1.7, 5.37, 3.0, 0.45, "🏁  结束", GRAY, WHITE, 12, True)

# 旁注:记录完整轨迹
txt(5.20, 2.6, 1.05, 2.4, "↑\n先记录\n完整\n执行\n轨迹\n构建\n因果图", GRAY, 9, False)

# ===== Column 2: 根因分析流程 =====
txt(6.5, 1.28, 4.0, 0.4, "② 根因分析流程", NAVY, 15, bold=True)

box(6.70, 1.82, 2.60, 0.40, "1️⃣  用户输入问题",      LIGHT_GRAY, NAVY,  11, True)
arrow(7.70, 2.27, 0.60, 0.22, BLUE)

box(6.70, 2.54, 2.60, 0.40, "2️⃣  构建因果图",        LIGHT_GRAY, NAVY,  11, True)
arrow(7.70, 2.99, 0.60, 0.22, BLUE)

box(6.70, 3.26, 2.60, 0.40, "3️⃣  描述问题现象",      LIGHT_GRAY, NAVY,  11, True)
arrow(7.70, 3.71, 0.60, 0.22, BLUE)

box(6.70, 3.98, 2.60, 0.40, "4️⃣  触发回溯  ←",      BLUE,       WHITE, 12, True)
arrow(7.70, 4.43, 0.60, 0.22, BLUE)

# 分支:5Why + CAE
box(6.50, 4.70, 1.30, 0.36, "5Why", YELLOW, WHITE, 10, True)
box(8.20, 4.70, 1.30, 0.36, "CAE",  PURPLE, WHITE, 10, True)

arrow(7.70, 5.11, 0.60, 0.20, BLUE)
box(6.70, 5.36, 2.60, 0.40, "5️⃣  推理定位根因",      LIGHT_GRAY, NAVY,  11, True)
arrow(7.70, 5.81, 0.60, 0.22, BLUE)
box(6.70, 6.08, 2.60, 0.40, "6️⃣  输出根因结论",      GREEN,      WHITE, 11, True)

# ===== Column 3: 推理方法详解 =====
txt(9.5, 1.28, 3.5, 0.4, "③ 推理方法详解", NAVY, 15, bold=True)

# 5Why box
b1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(9.5), Inches(1.82), Inches(3.5), Inches(2.0))
b1.fill.solid(); b1.fill.fore_color.rgb = WHITE
b1.line.color.rgb = YELLOW; b1.line.width = Pt(2)
tf = b1.text_frame; tf.word_wrap = True
tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15); tf.margin_top = Inches(0.10)
tf.text = "5Why 分析法"
tf.paragraphs[0].font.size = Pt(14); tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = YELLOW
p = tf.add_paragraph(); p.text = "\n连续追问 5 次 \"为什么\""
p.font.size = Pt(11); p.font.color.rgb = DARK
p = tf.add_paragraph(); p.text = "由果 → 因, 逐层深入"
p.font.size = Pt(11); p.font.color.rgb = GRAY; p.font.italic = True
p = tf.add_paragraph(); p.text = "\n例: 为什么失败? → 因 X\n   为什么 X? → 因 Y\n   ... 直到根因"
p.font.size = Pt(10); p.font.color.rgb = GRAY

# CAE box
b2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(9.5), Inches(3.97), Inches(3.5), Inches(2.0))
b2.fill.solid(); b2.fill.fore_color.rgb = WHITE
b2.line.color.rgb = PURPLE; b2.line.width = Pt(2)
tf = b2.text_frame; tf.word_wrap = True
tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15); tf.margin_top = Inches(0.10)
tf.text = "CAE 模型"
tf.paragraphs[0].font.size = Pt(14); tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = PURPLE
p = tf.add_paragraph(); p.text = "\nClaim      主张 / 结论"
p.font.size = Pt(11); p.font.color.rgb = DARK
p = tf.add_paragraph(); p.text = "Argument   推理 / 论据"
p.font.size = Pt(11); p.font.color.rgb = DARK
p = tf.add_paragraph(); p.text = "Evidence   证据 / 数据"
p.font.size = Pt(11); p.font.color.rgb = DARK

# ===== Footer =====
fb = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(0.5), Inches(6.85), Inches(12.33), Inches(0.50))
fb.fill.solid(); fb.fill.fore_color.rgb = NAVY; fb.line.fill.background()
tf = fb.text_frame
tf.margin_left = Inches(0.15)
tf.text = "💡  核心思路: 先记录完整的执行轨迹(因果图)  →  用户描述问题  →  在图中反向回溯  +  推理方法定位根因"
tf.paragraphs[0].font.size = Pt(13); tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = WHITE; tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# ===== Save =====
out = "agent_root_cause_analysis.pptx"
prs.save(out)
print(f"[OK] Generated: {out}")